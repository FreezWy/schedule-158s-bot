
import os
import re
import logging
from datetime import datetime
from zoneinfo import ZoneInfo
from urllib.parse import urljoin

import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook
from io import BytesIO

SITE = "https://medcol-ptz.ru/student/studies/"
TOKEN = os.environ["TELEGRAM_BOT_TOKEN"]
CHAT_ID = os.environ["TELEGRAM_CHAT_ID"]
TZ = ZoneInfo("Europe/Moscow")

logging.basicConfig(level=logging.INFO)
log = logging.getLogger(__name__)


def get_latest_schedule():
    response = requests.get(
        SITE, timeout=30,
        headers={"User-Agent": "Mozilla/5.0"}
    )
    response.raise_for_status()
    soup = BeautifulSoup(response.text, "html.parser")

    candidates = []
    for a in soup.find_all("a", href=True):
        name = " ".join(a.get_text(" ", strip=True).split())
        href = urljoin(SITE, a["href"])
        text = name + " " + href

        if not href.lower().split("?")[0].endswith((".xlsx", ".xls")):
            continue
        if not re.search(r"156\s*с\s*[-–]\s*158\s*с", text, re.I):
            continue

        dates = re.findall(
            r"(\d{2}\.\d{2}\.\d{4})", name
        )
        if dates:
            try:
                start = datetime.strptime(dates[-2], "%d.%m.%Y")
                end = datetime.strptime(dates[-1], "%d.%m.%Y")
            except (ValueError, IndexError):
                continue
        else:
            # Не выбираем файл без периода из названия.
            continue

        candidates.append((start, end, name, href))

    if not candidates:
        raise RuntimeError(
            "Не найден Excel для групп 156с–158с "
            "с датами в названии."
        )

    # Выбираем публикацию с самым поздним периодом.
    _, _, name, url = max(candidates, key=lambda x: (x[0], x[1]))
    log.info("Выбран файл: %s", name)

    file_response = requests.get(url, timeout=60)
    file_response.raise_for_status()

    if url.lower().split("?")[0].endswith(".xls"):
        raise RuntimeError(
            "Найден старый формат .xls. "
            "Нужна поддержка xlrd или преобразование в .xlsx."
        )

    return name, load_workbook(
        BytesIO(file_response.content),
        data_only=True,
        read_only=True
    )


def cell_text(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%d.%m.%Y %H:%M")
    return str(value).strip()


def extract_group(workbook):
    # Сначала ищем отдельный лист для 158с.
    matching_sheets = [
        ws for ws in workbook.worksheets
        if re.search(r"(?<!\w)158\s*с(?!\w)", ws.title, re.I)
    ]

    # Если отдельного листа нет, ищем упоминание 158с
    # в содержимом листов.
    if not matching_sheets:
        for ws in workbook.worksheets:
            for row in ws.iter_rows(values_only=True):
                if any(
                    re.search(
                        r"(?<!\w)158\s*с(?!\w)",
                        cell_text(v), re.I
                    )
                    for v in row
                ):
                    matching_sheets.append(ws)
                    break

    if not matching_sheets:
        raise RuntimeError(
            "Не удалось определить лист или блок группы 158с. "
            "Нужно проверить структуру Excel."
        )

    lines = []
    for ws in matching_sheets:
        lines.append(f"📚 Лист: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            values = [cell_text(v) for v in row]
            while values and not values[-1]:
                values.pop()
            if any(values):
                lines.append(" | ".join(values))

    return "\n".join(lines)


def send_telegram(text):
    url = f"https://api.telegram.org/bot{TOKEN}/sendMessage"

    # Telegram ограничивает сообщение 4096 символами.
    # Разбиваем по строкам, не разрезая строку посередине.
    chunks = []
    current = ""
    for line in text.splitlines():
        if len(current) + len(line) + 1 > 3800:
            if current:
                chunks.append(current)
            current = line
        else:
            current = f"{current}\n{line}".strip()

    if current:
        chunks.append(current)

    for chunk in chunks:
        response = requests.post(
            url,
            json={
                "chat_id": CHAT_ID,
                "text": chunk,
                "disable_web_page_preview": True
            },
            timeout=30
        )
        response.raise_for_status()
        data = response.json()
        if not data.get("ok"):
            raise RuntimeError(f"Telegram API: {data}")


def main():
    filename, workbook = get_latest_schedule()
    try:
        schedule = extract_group(workbook)
    finally:
        workbook.close()

    if not schedule.strip():
        raise RuntimeError("Расписание пустое.")

    today = datetime.now(TZ).strftime("%d.%m.%Y")
    message = (
        f"🩺 РАСПИСАНИЕ ГРУППЫ 158с\n"
        f"📅 Проверено: {today}\n"
        f"📄 Источник: {filename}\n\n"
        f"{schedule}\n\n"
        f"🌐 Официальный сайт: {SITE}"
    )
    send_telegram(message)
    log.info("Расписание отправлено.")


if __name__ == "__main__":
    main()
